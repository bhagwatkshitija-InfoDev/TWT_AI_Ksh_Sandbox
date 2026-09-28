"""DOCX parsing using python-docx."""

from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

from claude_mcp_docqa_agent.errors.exceptions import DOCXParsingError
from claude_mcp_docqa_agent.utils.logger import get_logger

logger = get_logger(__name__)


class DOCXParser:
    """Parse and extract content from DOCX files."""

    def __init__(self, file_path: str) -> None:
        """Initialize DOCX parser.

        Args:
            file_path: Path to DOCX file

        Raises:
            DOCXParsingError: If file is not a valid DOCX
        """
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise DOCXParsingError(f"DOCX file not found: {file_path}")

        if self.file_path.suffix.lower() != ".docx":
            raise DOCXParsingError(f"File is not a DOCX: {file_path}")

        try:
            self.doc = Document(str(self.file_path))
            logger.info(f"Initialized DOCX parser for: {self.file_path}")
        except Exception as e:
            raise DOCXParsingError(f"Failed to open DOCX file: {e}")

    def extract_full_text(self) -> str:
        """Extract all text from DOCX.

        Returns:
            Combined text from all paragraphs
        """
        text_parts = [paragraph.text for paragraph in self.doc.paragraphs]
        full_text = "\n".join(text_parts)
        logger.info(f"Extracted text from {len(self.doc.paragraphs)} paragraphs")
        return full_text

    def extract_paragraphs(self) -> list[dict[str, Any]]:
        """Extract paragraphs with formatting information.

        Returns:
            List of paragraph dictionaries with text, style, and formatting
        """
        paragraphs = []

        for idx, para in enumerate(self.doc.paragraphs):
            # Extract outline level from style hierarchy
            level = 0
            if para.style and "Heading" in para.style.name:
                try:
                    level = int(para.style.name.split()[-1]) - 1
                except (ValueError, IndexError):
                    level = 0

            para_info = {
                "index": idx,
                "text": para.text,
                "style": para.style.name if para.style else "Normal",
                "alignment": str(para.alignment),
                "level": level,
                "runs": [],
            }

            # Extract run-level formatting
            for run in para.runs:
                run_info = {
                    "text": run.text,
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": run.underline,
                    "font_name": run.font.name,
                    "font_size": run.font.size.pt if run.font.size else None,
                    "color": run.font.color.rgb if run.font.color else None,
                }
                para_info["runs"].append(run_info)

            paragraphs.append(para_info)

        return paragraphs

    def extract_tables(self) -> list[dict[str, Any]]:
        """Extract tables with structure information.

        Returns:
            List of table dictionaries with rows and cells
        """
        tables = []

        for table_idx, table in enumerate(self.doc.tables):
            table_data = {
                "index": table_idx,
                "rows": [],
                "num_rows": len(table.rows),
                "num_cols": len(table.columns),
            }

            for row_idx, row in enumerate(table.rows):
                row_data = []
                for cell_idx, cell in enumerate(row.cells):
                    cell_data = {
                        "text": cell.text,
                        "paragraphs": len(cell.paragraphs),
                        "width": cell.width,
                    }
                    row_data.append(cell_data)
                table_data["rows"].append(row_data)

            tables.append(table_data)

        logger.info(f"Extracted {len(tables)} tables")
        return tables

    def extract_headings(self) -> list[dict[str, Any]]:
        """Extract heading structure.

        Returns:
            List of headings with level and text
        """
        headings = []

        for para in self.doc.paragraphs:
            if para.style and "Heading" in para.style.name:
                try:
                    level = int(para.style.name.split()[-1])
                except (ValueError, IndexError):
                    level = 0

                headings.append(
                    {
                        "text": para.text,
                        "style": para.style.name,
                        "level": level,
                    }
                )

        logger.info(f"Extracted {len(headings)} headings")
        return headings

    def extract_metadata(self) -> dict[str, Any]:
        """Extract document metadata.

        Returns:
            Dictionary with document properties
        """
        core_props = self.doc.core_properties

        return {
            "title": core_props.title,
            "author": core_props.author,
            "subject": core_props.subject,
            "keywords": core_props.keywords,
            "comments": core_props.comments,
            "created": str(core_props.created),
            "modified": str(core_props.modified),
            "num_paragraphs": len(self.doc.paragraphs),
            "num_tables": len(self.doc.tables),
            "num_sections": len(self.doc.sections),
            "file_size": self.file_path.stat().st_size,
        }

    def extract_font_usage(self) -> dict[str, int]:
        """Extract font usage statistics.

        Returns:
            Dictionary with font names and their frequency
        """
        font_usage = {}

        for para in self.doc.paragraphs:
            for run in para.runs:
                if run.font.name:
                    font_usage[run.font.name] = font_usage.get(run.font.name, 0) + 1

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.font.name:
                                font_usage[run.font.name] = (
                                    font_usage.get(run.font.name, 0) + 1
                                )

        logger.info(f"Found {len(font_usage)} unique fonts")
        return font_usage

    def extract_structured_content(self) -> dict[str, Any]:
        """Extract complete structured content.

        Returns:
            Comprehensive document structure dictionary
        """
        return {
            "file_path": str(self.file_path),
            "metadata": self.extract_metadata(),
            "paragraphs": self.extract_paragraphs(),
            "tables": self.extract_tables(),
            "headings": self.extract_headings(),
            "font_usage": self.extract_font_usage(),
            "full_text": self.extract_full_text(),
        }
