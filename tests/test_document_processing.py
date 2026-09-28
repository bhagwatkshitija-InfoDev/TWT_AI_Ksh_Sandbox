"""Unit tests for document processing modules."""

import json
import tempfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from claude_mcp_docqa_agent.document_processing.content_processor import ContentProcessor
from claude_mcp_docqa_agent.document_processing.docx_parser import DOCXParser
from claude_mcp_docqa_agent.document_processing.metadata_extractor import MetadataExtractor
from claude_mcp_docqa_agent.document_processing.pdf_extractor import PDFExtractor
from claude_mcp_docqa_agent.errors.exceptions import (
    DOCXParsingError,
    DocumentProcessingError,
    InvalidDocumentError,
)
from claude_mcp_docqa_agent.utils.validators import validate_document_path


@pytest.fixture
def temp_docx():
    """Create a temporary DOCX file for testing."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc = Document()

        # Add paragraphs with different styles
        para1 = doc.add_paragraph("Test Document", style="Heading 1")
        para2 = doc.add_paragraph("This is a test paragraph.", style="Normal")
        para3 = doc.add_paragraph(
            "This is another paragraph with some formatting.",
            style="Normal",
        )

        # Add run with formatting
        para4 = doc.add_paragraph()
        para4.add_run("Bold text ").bold = True
        para4.add_run("Italic text ").italic = True

        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Data 1"
        table.cell(1, 1).text = "Data 2"

        # Add core properties
        doc.core_properties.title = "Test Document"
        doc.core_properties.author = "Test Author"

        doc.save(f.name)
        yield f.name

    Path(f.name).unlink()


class TestDocxParser:
    """Test DOCX parsing functionality."""

    def test_parser_initialization(self, temp_docx):
        """Test DOCX parser initialization."""
        parser = DOCXParser(temp_docx)
        assert parser.file_path == Path(temp_docx)

    def test_parser_invalid_file(self):
        """Test parser with invalid file."""
        with pytest.raises(DOCXParsingError):
            DOCXParser("/nonexistent/file.docx")

    def test_parser_wrong_format(self, tmp_path):
        """Test parser with wrong file format."""
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("test")

        with pytest.raises(DOCXParsingError):
            DOCXParser(str(txt_file))

    def test_extract_full_text(self, temp_docx):
        """Test full text extraction."""
        parser = DOCXParser(temp_docx)
        text = parser.extract_full_text()

        assert isinstance(text, str)
        assert len(text) > 0
        assert "Test Document" in text
        assert "test paragraph" in text

    def test_extract_paragraphs(self, temp_docx):
        """Test paragraph extraction with formatting."""
        parser = DOCXParser(temp_docx)
        paragraphs = parser.extract_paragraphs()

        assert isinstance(paragraphs, list)
        assert len(paragraphs) > 0

        # Check first paragraph structure
        first_para = paragraphs[0]
        assert "text" in first_para
        assert "style" in first_para
        assert "runs" in first_para

    def test_extract_tables(self, temp_docx):
        """Test table extraction."""
        parser = DOCXParser(temp_docx)
        tables = parser.extract_tables()

        assert isinstance(tables, list)
        assert len(tables) > 0

        first_table = tables[0]
        assert "rows" in first_table
        assert first_table["num_rows"] == 2
        assert first_table["num_cols"] == 2

    def test_extract_headings(self, temp_docx):
        """Test heading extraction."""
        parser = DOCXParser(temp_docx)
        headings = parser.extract_headings()

        assert isinstance(headings, list)
        assert len(headings) > 0
        assert "Test Document" in headings[0]["text"]

    def test_extract_metadata(self, temp_docx):
        """Test metadata extraction."""
        parser = DOCXParser(temp_docx)
        metadata = parser.extract_metadata()

        assert "title" in metadata
        assert "author" in metadata
        assert "num_paragraphs" in metadata
        assert "num_tables" in metadata
        assert metadata["title"] == "Test Document"
        assert metadata["author"] == "Test Author"

    def test_extract_font_usage(self, temp_docx):
        """Test font usage extraction."""
        parser = DOCXParser(temp_docx)
        fonts = parser.extract_font_usage()

        assert isinstance(fonts, dict)

    def test_extract_structured_content(self, temp_docx):
        """Test complete structured content extraction."""
        parser = DOCXParser(temp_docx)
        content = parser.extract_structured_content()

        assert "file_path" in content
        assert "metadata" in content
        assert "paragraphs" in content
        assert "tables" in content
        assert "headings" in content
        assert "font_usage" in content
        assert "full_text" in content


class TestValidators:
    """Test input validators."""

    def test_validate_document_path_valid(self, temp_docx):
        """Test validation of valid document path."""
        path = validate_document_path(temp_docx)
        assert isinstance(path, Path)
        assert path.exists()

    def test_validate_document_path_not_found(self):
        """Test validation of non-existent file."""
        with pytest.raises(InvalidDocumentError):
            validate_document_path("/nonexistent/file.docx")

    def test_validate_document_path_invalid_format(self, tmp_path):
        """Test validation of invalid file format."""
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("test")

        with pytest.raises(InvalidDocumentError):
            validate_document_path(str(txt_file))


class TestMetadataExtractor:
    """Test metadata extraction."""

    def test_extractor_initialization(self, temp_docx):
        """Test metadata extractor initialization."""
        extractor = MetadataExtractor(temp_docx)
        assert extractor.file_path == Path(temp_docx)

    def test_extract_docx_metadata(self, temp_docx):
        """Test DOCX metadata extraction."""
        extractor = MetadataExtractor(temp_docx)
        metadata = extractor.extract_metadata()

        assert metadata["file_type"] == "docx"
        assert "filename" in metadata
        assert "file_size" in metadata
        assert "properties" in metadata

    def test_get_text_sample(self, temp_docx):
        """Test text sample extraction."""
        extractor = MetadataExtractor(temp_docx)
        sample = extractor.get_text_sample(length=50)

        assert isinstance(sample, str)
        assert len(sample) <= 50


class TestContentProcessor:
    """Test content processing."""

    def test_processor_initialization(self, temp_docx):
        """Test content processor initialization."""
        processor = ContentProcessor(temp_docx)
        assert processor.file_path == temp_docx

    def test_process_docx_document(self, temp_docx):
        """Test processing DOCX document."""
        processor = ContentProcessor(temp_docx)
        result = processor.process_document()

        assert result["file_type"] == "docx"
        assert "language" in result
        assert "language_detection" in result
        assert "content" in result
        assert "full_text" in result

    def test_get_text_for_analysis(self, temp_docx):
        """Test text extraction for analysis."""
        processor = ContentProcessor(temp_docx)
        text = processor.get_text_for_analysis()

        assert isinstance(text, str)
        assert len(text) > 0

    def test_get_page_count_docx(self, temp_docx):
        """Test page count for DOCX."""
        processor = ContentProcessor(temp_docx)
        count = processor.get_page_count()

        assert isinstance(count, int)
        assert count > 0
