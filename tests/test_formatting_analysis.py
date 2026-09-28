"""Unit tests for formatting analysis modules."""

import tempfile
from pathlib import Path

import pytest
from docx import Document

from claude_mcp_docqa_agent.analysis.cross_ref_checker import CrossReferenceChecker
from claude_mcp_docqa_agent.analysis.docx_analyzers import (
    DOCXFontAnalyzer,
    DOCXHeadingAnalyzer,
    DOCXListAnalyzer,
    DOCXTableAnalyzer,
)
from claude_mcp_docqa_agent.analysis.formatting_analyzer import (
    FormattingAnalyzer,
    FormattingIssue,
)
from claude_mcp_docqa_agent.analysis.image_analyzer import ImageAnalyzer
from claude_mcp_docqa_agent.analysis.pdf_analyzers import (
    PDFFontAnalyzer,
    PDFStructureAnalyzer,
    PDFTableAnalyzer,
)
from claude_mcp_docqa_agent.document_processing.content_processor import (
    ContentProcessor,
)


@pytest.fixture
def temp_docx_with_issues():
    """Create a DOCX with formatting issues for testing."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc = Document()

        # Add H1
        doc.add_paragraph("Main Title", style="Heading 1")

        # Add H3 (skip H2 - should flag hierarchy issue)
        doc.add_paragraph("Subsection", style="Heading 3")

        # Add multiple fonts (should flag font consistency)
        para = doc.add_paragraph()
        run1 = para.add_run("Arial text ")
        run1.font.name = "Arial"
        run2 = para.add_run("Courier text ")
        run2.font.name = "Courier"
        run3 = para.add_run("Times text")
        run3.font.name = "Times New Roman"

        # Add table
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Col1"
        table.cell(0, 1).text = "Col2"
        table.cell(0, 2).text = "Col3"
        table.cell(1, 0).text = "Data1"
        table.cell(1, 1).text = "Data2"
        table.cell(1, 2).text = "Data3"

        # Add another H1 (multiple H1s - should flag)
        doc.add_paragraph("Another Title", style="Heading 1")

        doc.add_paragraph("Normal text.")

        doc.save(f.name)
        yield f.name

    Path(f.name).unlink()


@pytest.fixture
def sample_document_content():
    """Sample document content structure for testing."""
    return {
        "file_type": "docx",
        "language": "de",
        "full_text": "Main Title Subsection text content. See Section 1. Refer to Table 1.",
        "content": {
            "metadata": {
                "title": "Test Document",
                "num_paragraphs": 5,
                "num_tables": 1,
            },
            "paragraphs": [
                {"text": "Main Title", "style": "Heading 1", "level": 0, "runs": []},
                {
                    "text": "Subsection",
                    "style": "Heading 3",
                    "level": 2,
                    "runs": [],
                },
                {
                    "text": "Some normal text content.",
                    "style": "Normal",
                    "level": 0,
                    "runs": [],
                },
            ],
            "headings": [
                {"text": "Main Title", "style": "Heading 1", "level": 1},
                {
                    "text": "Subsection",
                    "style": "Heading 3",
                    "level": 3,
                },
            ],
            "tables": [
                {
                    "index": 0,
                    "num_rows": 2,
                    "num_cols": 3,
                    "rows": [
                        [{"text": "Col1"}, {"text": "Col2"}, {"text": "Col3"}],
                        [
                            {"text": "Data1"},
                            {"text": "Data2"},
                            {"text": "Data3"},
                        ],
                    ],
                }
            ],
            "font_usage": {
                "Arial": 100,
                "Courier": 50,
                "Times New Roman": 30,
                "Calibri": 20,
            },
        },
    }


class TestFormattingIssue:
    """Test FormattingIssue dataclass."""

    def test_issue_creation(self):
        """Test creating a formatting issue."""
        issue = FormattingIssue(
            id="test-1",
            page_number=1,
            issue_type="font",
            severity="warning",
            location_description="Header",
            issue_description="Font issue",
            recommended_fix="Use standard font",
        )

        assert issue.id == "test-1"
        assert issue.page_number == 1
        assert issue.severity == "warning"

    def test_issue_to_dict(self):
        """Test converting issue to dictionary."""
        issue = FormattingIssue(
            id="test-2",
            page_number=2,
            issue_type="heading",
            severity="info",
            location_description="Title",
            issue_description="Heading issue",
            recommended_fix="Fix heading",
        )

        issue_dict = issue.to_dict()
        assert isinstance(issue_dict, dict)
        assert issue_dict["id"] == "test-2"
        assert issue_dict["severity"] == "info"


class TestDOCXFontAnalyzer:
    """Test DOCX font analysis."""

    def test_font_analyzer_initialization(self, sample_document_content):
        """Test initializing font analyzer."""
        analyzer = DOCXFontAnalyzer(sample_document_content["content"])
        assert analyzer is not None

    def test_font_consistency_check(self, sample_document_content):
        """Test font consistency checking."""
        analyzer = DOCXFontAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        # Should find issue with 4 fonts (max is 3)
        assert len(issues) > 0
        assert any(i.issue_type == "font_variety" for i in issues)

    def test_get_font_list(self, sample_document_content):
        """Test getting sorted font list."""
        analyzer = DOCXFontAnalyzer(sample_document_content["content"])
        fonts = analyzer.get_font_list()

        assert len(fonts) == 4
        assert fonts[0][1] > fonts[-1][1]  # Sorted by usage


class TestDOCXHeadingAnalyzer:
    """Test DOCX heading analysis."""

    def test_heading_analyzer_initialization(self, sample_document_content):
        """Test initializing heading analyzer."""
        analyzer = DOCXHeadingAnalyzer(sample_document_content["content"])
        assert analyzer is not None

    def test_heading_hierarchy_check(self, sample_document_content):
        """Test heading hierarchy checking."""
        analyzer = DOCXHeadingAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        # Should find issue with H1 -> H3 skipping H2
        hierarchy_issues = [i for i in issues if i.issue_type == "heading_hierarchy"]
        assert len(hierarchy_issues) > 0

    def test_heading_tree(self, sample_document_content):
        """Test getting heading hierarchy tree."""
        analyzer = DOCXHeadingAnalyzer(sample_document_content["content"])
        tree = analyzer.get_heading_tree()

        assert len(tree) == 2
        assert tree[0]["level"] == 1
        assert tree[1]["level"] == 3


class TestDOCXTableAnalyzer:
    """Test DOCX table analysis."""

    def test_table_analyzer_initialization(self, sample_document_content):
        """Test initializing table analyzer."""
        analyzer = DOCXTableAnalyzer(sample_document_content["content"])
        assert analyzer is not None

    def test_table_consistency_check(self, sample_document_content):
        """Test table consistency checking."""
        analyzer = DOCXTableAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        # Valid table should not generate issues
        assert len(issues) == 0

    def test_table_stats(self, sample_document_content):
        """Test getting table statistics."""
        analyzer = DOCXTableAnalyzer(sample_document_content["content"])
        stats = analyzer.get_table_stats()

        assert stats["total_tables"] == 1
        assert stats["avg_rows"] == 2
        assert stats["avg_cols"] == 3


class TestDOCXListAnalyzer:
    """Test DOCX list analysis."""

    def test_list_analyzer_initialization(self, sample_document_content):
        """Test initializing list analyzer."""
        analyzer = DOCXListAnalyzer(sample_document_content["content"])
        assert analyzer is not None

    def test_analyze_lists(self, sample_document_content):
        """Test list analysis."""
        analyzer = DOCXListAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        # Test fixture has minimal lists
        assert isinstance(issues, list)


class TestPDFFontAnalyzer:
    """Test PDF font analysis."""

    def test_pdf_font_analyzer_initialization(self, sample_document_content):
        """Test PDF font analyzer initialization."""
        sample_document_content["file_type"] = "pdf"
        analyzer = PDFFontAnalyzer(sample_document_content["content"])
        assert analyzer is not None

    def test_pdf_analyze(self, sample_document_content):
        """Test PDF analysis."""
        sample_document_content["file_type"] = "pdf"
        analyzer = PDFFontAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        assert isinstance(issues, list)


class TestPDFStructureAnalyzer:
    """Test PDF structure analysis."""

    def test_pdf_structure_analyzer(self, sample_document_content):
        """Test PDF structure analyzer."""
        sample_document_content["content"]["pages"] = [
            {"num_chars": 100, "num_lines": 10, "num_rects": 2, "tables": []},
            {"num_chars": 0, "num_lines": 0, "num_rects": 0, "tables": []},
        ]

        analyzer = PDFStructureAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        # Second page is blank
        blank_issues = [i for i in issues if i.issue_type == "page_structure"]
        assert len(blank_issues) > 0


class TestImageAnalyzer:
    """Test image analysis."""

    def test_image_analyzer_initialization(self, sample_document_content):
        """Test initializing image analyzer."""
        analyzer = ImageAnalyzer(sample_document_content["content"])
        assert analyzer is not None

    def test_analyze_images(self, sample_document_content):
        """Test image analysis."""
        analyzer = ImageAnalyzer(sample_document_content["content"])
        issues = analyzer.analyze()

        assert isinstance(issues, list)

    def test_get_image_count(self, sample_document_content):
        """Test getting image count."""
        analyzer = ImageAnalyzer(sample_document_content["content"])
        count = analyzer.get_image_count()

        assert isinstance(count, int)


class TestCrossReferenceChecker:
    """Test cross-reference analysis."""

    def test_cross_ref_checker_initialization(self, sample_document_content):
        """Test initializing cross-reference checker."""
        checker = CrossReferenceChecker(sample_document_content["content"])
        assert checker is not None

    def test_cross_ref_analysis(self, sample_document_content):
        """Test cross-reference analysis."""
        checker = CrossReferenceChecker(sample_document_content["content"])
        issues = checker.analyze()

        assert isinstance(issues, list)

    def test_reference_summary(self, sample_document_content):
        """Test getting reference summary."""
        checker = CrossReferenceChecker(sample_document_content["content"])
        summary = checker.get_reference_summary()

        assert "total_headings" in summary
        assert "total_tables" in summary
        assert "total_issues" in summary


class TestFormattingAnalyzer:
    """Test main formatting analyzer."""

    def test_analyzer_initialization(self, sample_document_content):
        """Test initializing formatting analyzer."""
        analyzer = FormattingAnalyzer(sample_document_content)
        assert analyzer is not None

    def test_analyze_all_docx(self, sample_document_content):
        """Test analyzing DOCX document."""
        analyzer = FormattingAnalyzer(sample_document_content)
        issues = analyzer.analyze_all()

        assert isinstance(issues, list)
        # Should find heading hierarchy issue
        assert any(i.issue_type == "heading_hierarchy" for i in issues)

    def test_get_issues_by_severity(self, sample_document_content):
        """Test filtering issues by severity."""
        analyzer = FormattingAnalyzer(sample_document_content)
        analyzer.analyze_all()

        warnings = analyzer.get_issues_by_severity("warning")
        assert isinstance(warnings, list)

    def test_get_issues_by_type(self, sample_document_content):
        """Test filtering issues by type."""
        analyzer = FormattingAnalyzer(sample_document_content)
        analyzer.analyze_all()

        heading_issues = analyzer.get_issues_by_type("heading_hierarchy")
        assert isinstance(heading_issues, list)

    def test_get_summary(self, sample_document_content):
        """Test getting issue summary."""
        analyzer = FormattingAnalyzer(sample_document_content)
        analyzer.analyze_all()

        summary = analyzer.get_summary()
        assert "total_issues" in summary
        assert "critical" in summary
        assert "warning" in summary
        assert "info" in summary
        assert "by_type" in summary


class TestFormattingAnalysisIntegration:
    """Integration tests for formatting analysis."""

    def test_end_to_end_docx_analysis(self, temp_docx_with_issues):
        """Test end-to-end DOCX analysis."""
        processor = ContentProcessor(temp_docx_with_issues)
        document_data = processor.process_document()

        analyzer = FormattingAnalyzer(document_data)
        issues = analyzer.analyze_all()

        # Should find multiple issues
        assert len(issues) > 0

        # Check issue types
        issue_types = set(i.issue_type for i in issues)
        assert any(
            t
            in issue_types
            for t in ["font_variety", "heading_hierarchy", "heading_consistency"]
        )

    def test_severity_distribution(self, sample_document_content):
        """Test severity distribution in issues."""
        analyzer = FormattingAnalyzer(sample_document_content)
        issues = analyzer.analyze_all()

        summary = analyzer.get_summary()

        total = (
            summary["critical"] + summary["warning"] + summary["info"]
        )
        assert total == summary["total_issues"]

    def test_page_filtering(self, sample_document_content):
        """Test page-based issue filtering."""
        analyzer = FormattingAnalyzer(sample_document_content)
        analyzer.analyze_all()

        page1_issues = analyzer.get_issues_by_page(1)
        assert isinstance(page1_issues, list)
