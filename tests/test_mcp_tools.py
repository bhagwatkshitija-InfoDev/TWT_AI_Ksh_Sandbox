"""Unit tests for MCP tools."""

import json
import tempfile
from pathlib import Path

import pytest
from docx import Document

from claude_mcp_docqa_agent.mcp_server.tools import (
    TOOLS,
    TOOL_HANDLERS,
    analyze_complete,
    analyze_formatting,
    analyze_language,
    get_document_summary,
    handle_tool_call,
    process_document,
    upload_document,
)


@pytest.fixture
def temp_docx():
    """Create a temporary DOCX file for testing."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc = Document()

        para1 = doc.add_paragraph("Test Document", style="Heading 1")
        para2 = doc.add_paragraph(
            "This is a test paragraph with sufficient content for language detection. "
            "The text needs to be at least 50 characters long for proper language detection to work. "
            "This paragraph contains more than enough text to meet that requirement.",
            style="Normal",
        )

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Data 1"
        table.cell(1, 1).text = "Data 2"

        doc.core_properties.title = "Test Document"
        doc.save(f.name)
        yield f.name

    try:
        Path(f.name).unlink()
    except PermissionError:
        pass  # File may still be in use


class TestToolDefinitions:
    """Test tool definitions."""

    def test_tools_defined(self):
        """Test that all tools are defined."""
        assert len(TOOLS) == 6

    def test_tool_names(self):
        """Test tool names."""
        tool_names = [t["name"] for t in TOOLS]

        expected_names = [
            "upload_document",
            "process_document",
            "analyze_formatting",
            "analyze_language",
            "analyze_complete",
            "get_document_summary",
        ]

        for name in expected_names:
            assert name in tool_names

    def test_tool_schemas(self):
        """Test that all tools have input schemas."""
        for tool in TOOLS:
            assert "inputSchema" in tool
            assert "type" in tool["inputSchema"]
            assert "properties" in tool["inputSchema"]
            assert "required" in tool["inputSchema"]

    def test_tool_handlers_registered(self):
        """Test that all tools have handlers."""
        tool_names = [t["name"] for t in TOOLS]

        for name in tool_names:
            assert name in TOOL_HANDLERS


class TestUploadDocumentTool:
    """Test upload_document tool."""

    def test_upload_valid_docx(self, temp_docx):
        """Test uploading a valid DOCX file."""
        result = upload_document(temp_docx)

        assert result["status"] == "success"
        assert "filename" in result
        assert result["file_type"] == "docx"
        assert "file_size" in result

    def test_upload_nonexistent_file(self):
        """Test uploading a nonexistent file."""
        result = upload_document("/nonexistent/file.docx")

        assert result["status"] == "error"
        assert "error" in result


class TestProcessDocumentTool:
    """Test process_document tool."""

    def test_process_valid_docx(self, temp_docx):
        """Test processing a valid DOCX file."""
        result = process_document(temp_docx)

        assert result["status"] == "success"
        assert result["file_type"] == "docx"
        assert "language" in result
        assert "text_sample" in result

    def test_process_nonexistent_file(self):
        """Test processing a nonexistent file."""
        result = process_document("/nonexistent/file.docx")

        assert result["status"] == "error"
        assert "error" in result


class TestAnalyzeFormattingTool:
    """Test analyze_formatting tool."""

    def test_analyze_formatting_valid_document(self, temp_docx):
        """Test analyzing formatting of a valid document."""
        result = analyze_formatting(temp_docx)

        assert result["status"] == "success"
        assert "total_formatting_issues" in result
        assert "critical" in result
        assert "warning" in result
        assert "info" in result

    def test_analyze_formatting_nonexistent_file(self):
        """Test analyzing formatting of a nonexistent file."""
        result = analyze_formatting("/nonexistent/file.docx")

        assert result["status"] == "error"


class TestAnalyzeLanguageTool:
    """Test analyze_language tool."""

    def test_analyze_language_valid_document(self, temp_docx):
        """Test analyzing language of a valid document."""
        result = analyze_language(temp_docx)

        assert result["status"] == "success"
        assert "language" in result
        assert "total_language_issues" in result

    def test_analyze_language_nonexistent_file(self):
        """Test analyzing language of a nonexistent file."""
        result = analyze_language("/nonexistent/file.docx")

        assert result["status"] == "error"


class TestAnalyzeCompleteTool:
    """Test analyze_complete tool."""

    def test_analyze_complete_valid_document(self, temp_docx):
        """Test complete analysis of a valid document."""
        result = analyze_complete(temp_docx)

        assert result["status"] == "success"
        assert "language" in result
        assert "file_type" in result
        assert "total_issues" in result
        assert "formatting_issues" in result
        assert "language_issues" in result
        assert "quality_score" in result
        assert 0 <= result["quality_score"] <= 100

    def test_analyze_complete_nonexistent_file(self):
        """Test complete analysis of a nonexistent file."""
        result = analyze_complete("/nonexistent/file.docx")

        assert result["status"] == "error"


class TestGetDocumentSummaryTool:
    """Test get_document_summary tool."""

    def test_get_summary_valid_document(self, temp_docx):
        """Test getting summary of a valid document."""
        result = get_document_summary(temp_docx)

        assert result["status"] == "success"
        assert "language" in result
        assert "file_type" in result
        assert "quality_score" in result
        assert "total_issues" in result
        assert "priority_issues" in result
        assert isinstance(result["priority_issues"], list)

    def test_get_summary_nonexistent_file(self):
        """Test getting summary of a nonexistent file."""
        result = get_document_summary("/nonexistent/file.docx")

        assert result["status"] == "error"


class TestToolDispatcher:
    """Test tool dispatcher."""

    def test_handle_upload_document_tool(self, temp_docx):
        """Test handling upload_document tool call."""
        result_str = handle_tool_call("upload_document", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "success"

    def test_handle_process_document_tool(self, temp_docx):
        """Test handling process_document tool call."""
        result_str = handle_tool_call("process_document", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "success"

    def test_handle_analyze_formatting_tool(self, temp_docx):
        """Test handling analyze_formatting tool call."""
        result_str = handle_tool_call("analyze_formatting", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "success"

    def test_handle_analyze_language_tool(self, temp_docx):
        """Test handling analyze_language tool call."""
        result_str = handle_tool_call("analyze_language", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "success"

    def test_handle_analyze_complete_tool(self, temp_docx):
        """Test handling analyze_complete tool call."""
        result_str = handle_tool_call("analyze_complete", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "success"

    def test_handle_get_summary_tool(self, temp_docx):
        """Test handling get_document_summary tool call."""
        result_str = handle_tool_call("get_document_summary", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "success"

    def test_handle_unknown_tool(self, temp_docx):
        """Test handling unknown tool."""
        result_str = handle_tool_call("unknown_tool", {"file_path": temp_docx})
        result = json.loads(result_str)

        assert result["status"] == "error"
        assert "Unknown tool" in result["error"]

    def test_handle_missing_parameter(self):
        """Test handling missing required parameter."""
        result_str = handle_tool_call("upload_document", {})
        result = json.loads(result_str)

        assert result["status"] == "error"


class TestMCPIntegration:
    """Integration tests for MCP tools."""

    def test_complete_workflow(self, temp_docx):
        """Test complete MCP workflow."""
        # Step 1: Upload
        upload_result = json.loads(
            handle_tool_call("upload_document", {"file_path": temp_docx})
        )
        assert upload_result["status"] == "success"

        # Step 2: Process
        process_result = json.loads(
            handle_tool_call("process_document", {"file_path": temp_docx})
        )
        assert process_result["status"] == "success"

        # Step 3: Analyze Formatting
        format_result = json.loads(
            handle_tool_call("analyze_formatting", {"file_path": temp_docx})
        )
        assert format_result["status"] == "success"

        # Step 4: Analyze Language
        lang_result = json.loads(
            handle_tool_call("analyze_language", {"file_path": temp_docx})
        )
        assert lang_result["status"] == "success"

        # Step 5: Complete Analysis
        complete_result = json.loads(
            handle_tool_call("analyze_complete", {"file_path": temp_docx})
        )
        assert complete_result["status"] == "success"

        # Step 6: Get Summary
        summary_result = json.loads(
            handle_tool_call("get_document_summary", {"file_path": temp_docx})
        )
        assert summary_result["status"] == "success"

    def test_json_serialization(self, temp_docx):
        """Test that all tool results are JSON serializable."""
        tool_calls = [
            ("upload_document", {"file_path": temp_docx}),
            ("process_document", {"file_path": temp_docx}),
            ("analyze_formatting", {"file_path": temp_docx}),
            ("analyze_language", {"file_path": temp_docx}),
            ("analyze_complete", {"file_path": temp_docx}),
            ("get_document_summary", {"file_path": temp_docx}),
        ]

        for tool_name, params in tool_calls:
            result_str = handle_tool_call(tool_name, params)

            # Should be valid JSON
            result = json.loads(result_str)
            assert isinstance(result, dict)
            assert "status" in result
